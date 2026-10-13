"""Écrit le classeur Excel désordonné à partir du dataset déjà défectueux.

Injecte en plus les défauts de structure (lignes de total, en-têtes
répétés, lignes vides, cellule fusionnée, colonne inutile) directement au
niveau du classeur, et calcule le numéro de ligne Excel final de chaque
ligne de données pour recaler la vérité terrain (voir
specs/01-generator.md et specs/02-ingestion.md pour le contrat de lecture
côté ingestion).
"""

import random
from pathlib import Path

from openpyxl import Workbook

from pipeline.schema import COLONNE_INUTILE_LABEL, TABLES as SHEETS

COLONNE_INUTILE_VALEURS = [
    None, None, None, "à vérifier", "RAS", "revu le mois dernier",
]

N_LIGNES_VIDES = 3
N_LIGNES_TOTAL = 2
N_ENTETES_REPETES = 2
N_BANNIERES = 1


def _evenements_structure(rng: random.Random, n_lignes: int) -> dict[int, str]:
    """Choisit des positions d'insertion (avant la ligne de données d'index
    0-based donné) pour chaque type de défaut de structure."""
    types = (
        ["ligne_vide"] * N_LIGNES_VIDES
        + ["ligne_total"] * N_LIGNES_TOTAL
        + ["entete_repete"] * N_ENTETES_REPETES
        + ["banniere"] * N_BANNIERES
    )
    positions = rng.sample(range(1, n_lignes), len(types))
    return dict(zip(positions, types))


def _ecrire_evenement(ws, type_evenement: str, excel_row: int, columns: list[str], labels: dict) -> None:
    n_cols = len(columns) + 1  # + colonne inutile
    if type_evenement == "ligne_vide":
        return
    if type_evenement == "ligne_total":
        ws.cell(row=excel_row, column=1, value="TOTAL")
        return
    if type_evenement == "entete_repete":
        for col_idx, colonne in enumerate(columns, start=1):
            ws.cell(row=excel_row, column=col_idx, value=labels[colonne])
        ws.cell(row=excel_row, column=len(columns) + 1, value=COLONNE_INUTILE_LABEL)
        return
    if type_evenement == "banniere":
        ws.cell(row=excel_row, column=1, value="--- Export mensuel - Service commercial ---")
        ws.merge_cells(start_row=excel_row, start_column=1, end_row=excel_row, end_column=min(4, n_cols))
        return
    raise ValueError(f"type d'évènement de structure inconnu : {type_evenement}")


def _ecrire_feuille(
    wb: Workbook,
    nom_feuille: str,
    rows: list[dict],
    rng: random.Random,
) -> tuple[dict[int, int], list[dict]]:
    """Retourne (mapping ligne_logique_1_based -> ligne_excel, defauts_structure)."""
    meta = SHEETS[nom_feuille]
    columns = meta["columns"]
    labels = meta["labels"]
    ws = wb.create_sheet(nom_feuille)

    for col_idx, colonne in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=labels[colonne])
    ws.cell(row=1, column=len(columns) + 1, value=COLONNE_INUTILE_LABEL)

    evenements = _evenements_structure(rng, len(rows))
    mapping: dict[int, int] = {}
    defauts_structure: list[dict] = []
    excel_row = 2

    for logique_idx, ligne in enumerate(rows):
        if logique_idx in evenements:
            type_evenement = evenements.pop(logique_idx)
            _ecrire_evenement(ws, type_evenement, excel_row, columns, labels)
            defauts_structure.append({
                "onglet": nom_feuille, "ligne_excel": excel_row,
                "type": type_evenement,
            })
            excel_row += 1

        for col_idx, colonne in enumerate(columns, start=1):
            ws.cell(row=excel_row, column=col_idx, value=ligne[colonne])
        ws.cell(
            row=excel_row, column=len(columns) + 1,
            value=rng.choice(COLONNE_INUTILE_VALEURS),
        )
        mapping[logique_idx + 1] = excel_row
        excel_row += 1

    defauts_structure.append({
        "onglet": nom_feuille, "ligne_excel": None, "type": "colonne_inutile",
    })

    return mapping, defauts_structure


def ecrire_classeur(dataset_desordonne: dict, chemin: Path, seed: int) -> tuple[dict, list[dict]]:
    """Écrit le classeur et retourne (mapping par onglet, défauts de structure)."""
    rng = random.Random(seed + 2)
    wb = Workbook()
    wb.remove(wb.active)

    mapping_global = {}
    structure_globale = []
    for nom_feuille, cle_dataset in [
        ("Clients", "clients"), ("Produits", "produits"), ("Commandes", "commandes"),
    ]:
        mapping, structure = _ecrire_feuille(wb, nom_feuille, dataset_desordonne[cle_dataset], rng)
        mapping_global[nom_feuille] = mapping
        structure_globale.extend(structure)

    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)
    return mapping_global, structure_globale
